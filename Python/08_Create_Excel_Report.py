import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.formatting.rule import ColorScaleRule
import os


# ============================================================
# CONFIGURATION
# ============================================================

input_file = "Data/processed/startup_funding_analytics.csv"
output_file = "Excel/Startup_Funding_Analysis.xlsx"

print("=" * 70)
print("CREATING EXCEL ANALYSIS REPORT")
print("=" * 70)


# ============================================================
# LOAD DATA
# ============================================================

df = pd.read_csv(input_file)

print(f"\nDataset loaded: {len(df):,} rows")


# ============================================================
# PREPARE ANALYSIS TABLES
# ============================================================

known_df = df.dropna(subset=["Amount_USD"]).copy()

total_funding = known_df["Amount_USD"].sum()
known_deals = len(known_df)
total_records = len(df)
unique_startups = df["Startup_Name"].nunique()
disclosure_rate = known_deals / total_records


# ------------------------------------------------------------
# KPI SUMMARY
# ------------------------------------------------------------

kpi = pd.DataFrame({
    "KPI": [
        "Total Records",
        "Unique Startups",
        "Known Funding Records",
        "Missing / Undisclosed Funding",
        "Funding Disclosure Rate",
        "Total Disclosed Funding",
        "Average Deal Size",
        "Median Deal Size",
        "Largest Deal",
        "Duplicate Records"
    ],
    "Value": [
        total_records,
        unique_startups,
        known_deals,
        df["Amount_USD"].isna().sum(),
        disclosure_rate,
        total_funding,
        known_df["Amount_USD"].mean(),
        known_df["Amount_USD"].median(),
        known_df["Amount_USD"].max(),
        df.duplicated().sum()
    ]
})


# ------------------------------------------------------------
# FUNDING BY YEAR
# ------------------------------------------------------------

yearly = (
    known_df
    .groupby("Funding_Year")
    .agg(
        Total_Funding=("Amount_USD", "sum"),
        Funding_Deals=("Amount_USD", "count"),
        Average_Deal=("Amount_USD", "mean"),
        Median_Deal=("Amount_USD", "median"),
        Unique_Startups=("Startup_Name", "nunique")
    )
    .reset_index()
)

yearly["Funding_Share_%"] = (
    yearly["Total_Funding"] / total_funding
)


# ------------------------------------------------------------
# INDUSTRY ANALYSIS
# ------------------------------------------------------------

industry = (
    known_df
    .groupby("Industry_Standardized")
    .agg(
        Total_Funding=("Amount_USD", "sum"),
        Funding_Deals=("Amount_USD", "count"),
        Average_Deal=("Amount_USD", "mean"),
        Unique_Startups=("Startup_Name", "nunique")
    )
    .reset_index()
    .sort_values("Total_Funding", ascending=False)
)

industry["Funding_Share_%"] = (
    industry["Total_Funding"] / total_funding
)


# ------------------------------------------------------------
# CITY ANALYSIS
# ------------------------------------------------------------

city = (
    known_df
    .groupby("City_Standardized")
    .agg(
        Total_Funding=("Amount_USD", "sum"),
        Funding_Deals=("Amount_USD", "count"),
        Unique_Startups=("Startup_Name", "nunique"),
        Average_Deal=("Amount_USD", "mean")
    )
    .reset_index()
    .sort_values("Total_Funding", ascending=False)
)

city["Funding_Share_%"] = (
    city["Total_Funding"] / total_funding
)


# ------------------------------------------------------------
# INVESTMENT TYPE
# ------------------------------------------------------------

investment = (
    known_df
    .groupby("Investment_Type_Standardized")
    .agg(
        Total_Funding=("Amount_USD", "sum"),
        Funding_Deals=("Amount_USD", "count"),
        Average_Deal=("Amount_USD", "mean")
    )
    .reset_index()
    .sort_values("Total_Funding", ascending=False)
)

investment["Funding_Share_%"] = (
    investment["Total_Funding"] / total_funding
)


# ------------------------------------------------------------
# FUNDING SIZE
# ------------------------------------------------------------

funding_size = (
    known_df
    .groupby("Funding_Size_Category")
    .agg(
        Total_Funding=("Amount_USD", "sum"),
        Funding_Deals=("Amount_USD", "count"),
        Average_Deal=("Amount_USD", "mean")
    )
    .reset_index()
)

funding_size["Funding_Share_%"] = (
    funding_size["Total_Funding"] / total_funding
)


# ------------------------------------------------------------
# TOP STARTUPS
# ------------------------------------------------------------

top_startups = (
    known_df
    .groupby("Startup_Name")
    .agg(
        Total_Funding=("Amount_USD", "sum"),
        Funding_Deals=("Amount_USD", "count")
    )
    .reset_index()
    .sort_values("Total_Funding", ascending=False)
)

top_startups["Funding_Share_%"] = (
    top_startups["Total_Funding"] / total_funding
)

top_startups = top_startups.head(100)


# ------------------------------------------------------------
# DATA QUALITY
# ------------------------------------------------------------

data_quality = pd.DataFrame({
    "Metric": [
        "Total Records",
        "Duplicate Records",
        "Known Funding Records",
        "Missing Funding Records",
        "Funding Disclosure Rate",
        "Unique Startups",
        "Unique Industries",
        "Unique Cities",
        "Unique Investors"
    ],
    "Value": [
        len(df),
        df.duplicated().sum(),
        known_deals,
        df["Amount_USD"].isna().sum(),
        disclosure_rate,
        df["Startup_Name"].nunique(),
        df["Industry_Standardized"].nunique(),
        df["City_Standardized"].nunique(),
        df["Investors_Name"].nunique()
    ]
})


# ------------------------------------------------------------
# DATA DICTIONARY
# ------------------------------------------------------------

data_dictionary = pd.DataFrame({
    "Column": df.columns,
    "Description": [
        "Original serial number",
        "Original funding date",
        "Startup name",
        "Original industry classification",
        "Startup sub-vertical",
        "Original city/location",
        "Investor or investor group name",
        "Original investment type",
        "Funding amount converted to numeric USD",
        "Additional remarks from source dataset",
        "Year extracted from funding date",
        "Month extracted from funding date",
        "Quarter extracted from funding date",
        "Funding size classification",
        "Standardized city grouping",
        "Standardized industry grouping",
        "Standardized investment type"
    ]
})


# ============================================================
# CREATE EXCEL FILE
# ============================================================

os.makedirs("Excel", exist_ok=True)

with pd.ExcelWriter(output_file, engine="openpyxl") as writer:

    df.to_excel(writer, sheet_name="Data", index=False)

    data_dictionary.to_excel(
        writer,
        sheet_name="Data Dictionary",
        index=False
    )

    kpi.to_excel(
        writer,
        sheet_name="KPI Summary",
        index=False
    )

    yearly.to_excel(
        writer,
        sheet_name="Funding by Year",
        index=False
    )

    industry.to_excel(
        writer,
        sheet_name="Industry Analysis",
        index=False
    )

    city.to_excel(
        writer,
        sheet_name="City Analysis",
        index=False
    )

    investment.to_excel(
        writer,
        sheet_name="Investment Type",
        index=False
    )

    funding_size.to_excel(
        writer,
        sheet_name="Funding Size",
        index=False
    )

    top_startups.to_excel(
        writer,
        sheet_name="Top Startups",
        index=False
    )

    data_quality.to_excel(
        writer,
        sheet_name="Data Quality",
        index=False
    )


# ============================================================
# FORMAT WORKBOOK
# ============================================================

wb = load_workbook(output_file)

# Professional header style
header_fill = PatternFill(
    fill_type="solid",
    fgColor="1F4E78"
)

header_font = Font(
    color="FFFFFF",
    bold=True
)

thin_border = Border(
    bottom=Side(
        style="thin",
        color="D9E1F2"
    )
)


for ws in wb.worksheets:

    # Freeze first row
    ws.freeze_panes = "A2"

    # Header formatting
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )

    # Header height
    ws.row_dimensions[1].height = 25

    # Auto-width columns
    for column_cells in ws.columns:

        max_length = 0
        column_letter = get_column_letter(
            column_cells[0].column
        )

        for cell in column_cells:

            if cell.value is not None:

                length = len(str(cell.value))

                if length > max_length:
                    max_length = length

        ws.column_dimensions[column_letter].width = min(
            max_length + 2,
            35
        )

    # Add borders
    for row in ws.iter_rows():

        for cell in row:

            cell.border = thin_border


# ============================================================
# NUMBER FORMATTING
# ============================================================

currency_sheets = [
    "KPI Summary",
    "Funding by Year",
    "Industry Analysis",
    "City Analysis",
    "Investment Type",
    "Funding Size",
    "Top Startups"
]

for sheet_name in currency_sheets:

    ws = wb[sheet_name]

    for row in ws.iter_rows():

        for cell in row:

            if isinstance(cell.value, (int, float)):

                # Large financial numbers
                if cell.column > 1:

                    cell.number_format = '#,##0.00'


# KPI specific formatting
ws = wb["KPI Summary"]

for row in range(2, ws.max_row + 1):

    kpi_name = ws.cell(row, 1).value

    if kpi_name == "Funding Disclosure Rate":

        ws.cell(row, 2).number_format = "0.00%"

    elif "Funding" in kpi_name or "Deal" in kpi_name:

        ws.cell(row, 2).number_format = '$#,##0'


# Percentage formatting
for sheet_name in [
    "Funding by Year",
    "Industry Analysis",
    "City Analysis",
    "Investment Type",
    "Funding Size",
    "Top Startups"
]:

    ws = wb[sheet_name]

    # Find percentage column
    for col in range(1, ws.max_column + 1):

        header = ws.cell(1, col).value

        if header == "Funding_Share_%":

            for row in range(2, ws.max_row + 1):

                ws.cell(row, col).number_format = "0.00%"


# ============================================================
# CONDITIONAL FORMATTING
# ============================================================

for sheet_name in [
    "Industry Analysis",
    "City Analysis",
    "Investment Type",
    "Top Startups"
]:

    ws = wb[sheet_name]

    # Find Total_Funding column
    funding_col = None

    for col in range(1, ws.max_column + 1):

        if ws.cell(1, col).value == "Total_Funding":

            funding_col = col
            break

    if funding_col:

        letter = get_column_letter(funding_col)

        ws.conditional_formatting.add(
            f"{letter}2:{letter}{ws.max_row}",
            ColorScaleRule(
                start_type="min",
                start_color="FFFFFF",
                mid_type="percentile",
                mid_value=50,
                mid_color="9DC3E6",
                end_type="max",
                end_color="1F4E78"
            )
        )


# ============================================================
# ADD CHARTS
# ============================================================

# Funding by Year chart
ws = wb["Funding by Year"]

chart = LineChart()

chart.title = "Funding Trend by Year"
chart.y_axis.title = "Total Funding"
chart.x_axis.title = "Year"

data = Reference(
    ws,
    min_col=2,
    min_row=1,
    max_row=ws.max_row
)

categories = Reference(
    ws,
    min_col=1,
    min_row=2,
    max_row=ws.max_row
)

chart.add_data(
    data,
    titles_from_data=True
)

chart.set_categories(categories)

chart.height = 7
chart.width = 14

ws.add_chart(chart, "H2")


# Industry chart
ws = wb["Industry Analysis"]

chart = BarChart()

chart.type = "bar"
chart.style = 10
chart.title = "Top Industries by Funding"
chart.y_axis.title = "Industry"
chart.x_axis.title = "Total Funding"

data = Reference(
    ws,
    min_col=2,
    min_row=1,
    max_row=min(ws.max_row, 11)
)

categories = Reference(
    ws,
    min_col=1,
    min_row=2,
    max_row=min(ws.max_row, 11)
)

chart.add_data(
    data,
    titles_from_data=True
)

chart.set_categories(categories)

chart.height = 8
chart.width = 15

ws.add_chart(chart, "H2")


# City chart
ws = wb["City Analysis"]

chart = BarChart()

chart.type = "bar"
chart.title = "Top Cities by Funding"
chart.y_axis.title = "City"
chart.x_axis.title = "Total Funding"

data = Reference(
    ws,
    min_col=2,
    min_row=1,
    max_row=min(ws.max_row, 11)
)

categories = Reference(
    ws,
    min_col=1,
    min_row=2,
    max_row=min(ws.max_row, 11)
)

chart.add_data(
    data,
    titles_from_data=True
)

chart.set_categories(categories)

chart.height = 8
chart.width = 15

ws.add_chart(chart, "H2")


# ============================================================
# SAVE
# ============================================================

wb.save(output_file)

print("\n" + "=" * 70)
print("EXCEL REPORT CREATED SUCCESSFULLY")
print("=" * 70)

print(f"\nSaved to:")
print(output_file)

print("\nSheets created:")

for sheet in wb.sheetnames:
    print(f"  - {sheet}")

print("\n" + "=" * 70)
